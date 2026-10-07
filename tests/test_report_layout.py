import sys
import unittest
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm


REPORTING = Path(__file__).resolve().parents[1] / "tools" / "reporting"
sys.path.insert(0, str(REPORTING))
from build_reports import add_table  # noqa: E402


class ReportLayoutTests(unittest.TestCase):
    def test_compact_table_uses_the_requested_width_instead_of_page_width(self):
        widths = [1.1, 1.0, 1.6, 1.4, 1.5]
        table = add_table(Document(), [["Puesto", "ET", "Línea km", "ET km", "Fitness"]], widths)
        grid_width = sum(int(column.get(qn("w:w"))) for column in table._tbl.tblGrid.gridCol_lst)

        self.assertFalse(table.autofit)
        self.assertAlmostEqual(grid_width, sum(Cm(width).twips for width in widths), delta=5)


if __name__ == "__main__":
    unittest.main()
