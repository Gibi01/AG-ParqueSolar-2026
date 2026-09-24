from __future__ import annotations

import json
from pathlib import Path
from typing import Iterable

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
ASSETS = ROOT / "tmp" / "report_assets"
OUT = ROOT / "Informes" / "docx"
OUT.mkdir(parents=True, exist_ok=True)
PROFILE = OUT / "Informe_Proyecto_Investigacion_Parques_Solares.docx"
ARTICLE = OUT / "Articulo_Catedra_Parques_Solares.docx"

BLACK = "222222"
GRAY = "E8E8E8"
MIDGRAY = "A6A6A6"


def load_run_summary() -> dict:
    metadata = json.loads((ROOT / "results" / "optimization_run.json").read_text(encoding="utf-8"))
    candidates = pd.read_csv(ROOT / "results" / "candidate_locations.csv")
    ranking = pd.read_csv(ROOT / "results" / "ranking.csv").sort_values("rank")
    valid = candidates[candidates.valid.astype(bool)].copy()
    valid["fitness"] = (
        metadata["fitness_weights"]["weight_solar"] * valid.solar_score
        + metadata["fitness_weights"]["weight_grid_distance"] * valid.grid_proximity_score
        + metadata["fitness_weights"]["weight_transformer_distance"] * valid.transformer_proximity_score
    )
    exhaustive = valid.nlargest(10, "fitness")
    ga_ids = ranking.grid_cell_id.astype(int).tolist()
    exhaustive_ids = exhaustive.grid_cell_id.astype(int).tolist()
    return {
        "metadata": metadata,
        "candidates": candidates,
        "ranking": ranking,
        "total": len(candidates),
        "valid": len(valid),
        "excluded": len(candidates) - len(valid),
        "excluded_pct": 100 * (len(candidates) - len(valid)) / len(candidates),
        "first": ranking.iloc[0],
        "solar_min": ranking.solar_annual_kwh_m2.min(),
        "solar_max": ranking.solar_annual_kwh_m2.max(),
        "zero_lines": int((ranking.distance_to_power_line_km == 0).sum()),
        "exhaustive_best": exhaustive.iloc[0],
        "exhaustive_ids": exhaustive_ids,
        "ga_ids": ga_ids,
        "exact_match": ga_ids == exhaustive_ids,
        "overlap": len(set(ga_ids) & set(exhaustive_ids)),
    }


RUN = load_run_summary()


REFERENCES = [
    "[1] J. H. Holland, Adaptation in Natural and Artificial Systems. University of Michigan Press, 1975.",
    "[2] D. E. Goldberg, Genetic Algorithms in Search, Optimization and Machine Learning. Addison-Wesley, 1989.",
    "[3] M. Srinivas y L. M. Patnaik, “Genetic algorithms: a survey”, Computer, vol. 27, n.º 6, pp. 17-26, 1994. https://doi.org/10.1109/2.294849",
    "[4] J. Malczewski, “GIS-based multicriteria decision analysis: a survey of the literature”, International Journal of Geographical Information Science, vol. 20, n.º 7, pp. 703-726, 2006. https://doi.org/10.1080/13658810600661508",
    "[5] M. Uyan, “GIS-based solar farms site selection using analytic hierarchy process in Karapinar region”, Renewable and Sustainable Energy Reviews, vol. 28, pp. 11-17, 2013. https://doi.org/10.1016/j.rser.2013.07.042",
    "[6] H. Z. Al Garni y A. Awasthi, “Solar PV power plant site selection using a GIS-AHP based approach with application in Saudi Arabia”, Applied Energy, vol. 206, pp. 1225-1240, 2017. https://doi.org/10.1016/j.apenergy.2017.10.024",
    "[7] J. Muñoz-Sabater et al., “ERA5-Land: a state-of-the-art global reanalysis dataset for land applications”, Earth System Science Data, vol. 13, pp. 4349-4383, 2021. https://doi.org/10.5194/essd-13-4349-2021",
    "[8] Copernicus Climate Change Service y ECMWF, ERA5-Land hourly Analysis Ready Cloud Optimised data on single levels from 1950 to present Product User Guide, 2026. https://confluence.ecmwf.int/spaces/CKB/pages/536218894/",
    "[9] Instituto Geográfico Nacional, Unidades Territoriales. Datos Argentina. https://datos.gob.ar/ar/dataset/ign-unidades-territoriales",
    "[10] Jefatura de Gabinete de Ministros, Localidades BAHRA. Datos Argentina. https://datos.gob.ar/dataset/jgm-servicio-normalizacion-datos-geograficos/archivo/jgm_8.21",
    "[11] Secretaría de Energía, Redes de distribución eléctrica. Datos Argentina. https://datos.gob.ar/dataset/energia-redes-distribucion-electrica",
    "[12] Secretaría de Energía, Transporte Eléctrico AT Estaciones Transformadoras. Datos Argentina. https://www.datos.gob.ar/dataset/energia-transporte-electrico-at-estaciones-transformadoras",
    "[13] OpenStreetMap contributors, Copyright and License. https://www.openstreetmap.org/copyright",
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_margins(cell, top=70, start=90, bottom=70, end=90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_font(run, name="Times New Roman", size=12, bold=None, italic=None, color=BLACK) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_keep(paragraph, keep_next=False, keep_lines=True) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    if keep_next:
        ppr.append(OxmlElement("w:keepNext"))
    if keep_lines:
        ppr.append(OxmlElement("w:keepLines"))


def configure_base(doc: Document, body_size=12, line=1.12, after=5) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(body_size)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = line
    normal.paragraph_format.space_after = Pt(after)
    for name, size, space_before, space_after in (
        ("Title", 22, 0, 16), ("Heading 1", 16, 14, 7),
        ("Heading 2", 13, 10, 5), ("Heading 3", 12, 7, 3),
    ):
        st = styles[name]
        st.font.name = "Times New Roman"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(BLACK)
        st.paragraph_format.space_before = Pt(space_before)
        st.paragraph_format.space_after = Pt(space_after)
        st.paragraph_format.keep_with_next = True
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)


def add_heading(doc: Document, text: str, level=1):
    p = doc.add_heading(text, level=level)
    set_keep(p, keep_next=True)
    return p


def add_p(doc: Document, text: str, *, bold_lead: str | None = None, style=None, italic=False, align=None):
    p = doc.add_paragraph(style=style)
    if bold_lead and text.startswith(bold_lead):
        a, b = text[:len(bold_lead)], text[len(bold_lead):]
        r = p.add_run(a)
        set_repeat_font(r, bold=True)
        r = p.add_run(b)
        set_repeat_font(r, italic=italic)
    else:
        r = p.add_run(text)
        set_repeat_font(r, italic=italic)
    if align is not None:
        p.alignment = align
    return p


def add_bullets(doc: Document, items: Iterable[str], level=0):
    for text in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.paragraph_format.space_after = Pt(3)
        set_repeat_font(p.add_run(text))


def add_numbered(doc: Document, items: Iterable[str]):
    for text in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        set_repeat_font(p.add_run(text))


def add_caption(doc: Document, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    set_repeat_font(p.add_run(text), size=10, italic=True)
    set_keep(p)
    return p


def add_picture(doc: Document, filename: str, width_cm: float, caption: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(ASSETS / filename), width=Cm(width_cm))
    set_keep(p, keep_next=True)
    add_caption(doc, caption)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths=None, font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    if widths:
        total_twips = int(sum(widths) * 567)
        tbl_pr = table._tbl.tblPr
        tbl_w = tbl_pr.find(qn("w:tblW"))
        if tbl_w is None:
            tbl_w = OxmlElement("w:tblW")
            tbl_pr.append(tbl_w)
        tbl_w.set(qn("w:w"), str(total_twips))
        tbl_w.set(qn("w:type"), "dxa")
        layout = tbl_pr.find(qn("w:tblLayout"))
        if layout is None:
            layout = OxmlElement("w:tblLayout")
            tbl_pr.append(layout)
        layout.set(qn("w:type"), "fixed")
        if sum(widths) <= 8:
            table.alignment = WD_TABLE_ALIGNMENT.LEFT
        for i, width in enumerate(widths):
            twips = str(int(width * 567))
            table.columns[i].width = Cm(width)
            table._tbl.tblGrid.gridCol_lst[i].set(qn("w:w"), twips)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, GRAY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_repeat_font(p.add_run(h), size=font_size, bold=True)
        if widths:
            cell.width = Cm(widths[i])
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            tc_w.set(qn("w:w"), str(int(widths[i] * 567)))
            tc_w.set(qn("w:type"), "dxa")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_margins(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            set_repeat_font(p.add_run(str(value)), size=font_size)
            if widths:
                cells[i].width = Cm(widths[i])
                tc_w = cells[i]._tc.get_or_add_tcPr().find(qn("w:tcW"))
                tc_w.set(qn("w:w"), str(int(widths[i] * 567)))
                tc_w.set(qn("w:type"), "dxa")
    return table


def profile_document() -> None:
    doc = Document()
    configure_base(doc, body_size=11.5, line=1.15, after=5)
    sec = doc.sections[0]

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(58)
    set_repeat_font(p.add_run("Universidad Tecnológica Nacional"), size=15, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_repeat_font(p.add_run("Facultad Regional Buenos Aires"), size=13, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(38)
    set_repeat_font(p.add_run("Optimización de ubicaciones de parques solares mediante algoritmos genéticos"), size=22, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    set_repeat_font(p.add_run("Documento guía de la investigación y concreción del modelo"), size=14, italic=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(35)
    set_repeat_font(p.add_run("Asignatura Algoritmos Genéticos"), size=12, bold=True)
    add_table(doc, ["Integrantes", "Legajos"], [["", ""], ["", ""]], widths=[8, 6], font_size=11)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(35)
    set_repeat_font(p.add_run("Ciclo lectivo 2026"), size=12)
    doc.add_page_break()

    add_heading(doc, "Índice de contenidos", 1)
    toc = [
        "1 Denominación del proyecto", "2 Situación problemática", "3 Problema de investigación",
        "4 Objetivos de la investigación", "5 Hipótesis y alcance", "6 Marco teórico",
        "7 Concreción del modelo", "8 Diseño experimental", "9 Resultados",
        "10 Discusión", "11 Conclusiones", "12 Líneas de continuidad", "Referencias", "Anexo técnico",
    ]
    for item in toc:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(.3)
        p.paragraph_format.space_after = Pt(4)
        set_repeat_font(p.add_run(item), size=11.5)

    add_heading(doc, "Resumen ejecutivo", 1)
    add_p(doc, f"Este proyecto desarrolla y verifica un prototipo en Python que clasifica celdas de una grilla sobre la provincia de Santa Fe para identificar ubicaciones potencialmente favorables para un parque solar fotovoltaico de dos hectáreas. El modelo integra radiación solar de ERA5-Land ARCO, distancia a líneas eléctricas, distancia a estaciones transformadoras y una exclusión urbana. La corrida analizada evaluó {RUN['total']} celdas de 9 km, excluyó {RUN['excluded']} por intersección con áreas urbanas aproximadas y optimizó las {RUN['valid']} restantes. La mejor solución hallada por el algoritmo fue la celda {int(RUN['first'].grid_cell_id)}, con fitness {RUN['first'].fitness:.4f} y radiación anual estimada de {RUN['first'].solar_annual_kwh_m2:.1f} kWh/m²/año. El TOP 10 genético compartió {RUN['overlap']} posiciones con el TOP 10 exhaustivo y no recuperó la mejor alternativa de referencia, la celda {int(RUN['exhaustive_best'].grid_cell_id)}. La salida constituye una herramienta de preselección y no una decisión de inversión.")

    add_heading(doc, "1 Denominación del proyecto", 1)
    add_p(doc, "Optimización de ubicaciones potenciales para parques solares fotovoltaicos en Santa Fe mediante algoritmos genéticos y análisis geoespacial.")

    add_heading(doc, "2 Situación problemática", 1)
    add_p(doc, "La selección preliminar de un sitio para generación fotovoltaica es un problema espacial de decisión multicriterio. Una ubicación puede recibir buena radiación y, al mismo tiempo, encontrarse lejos de la infraestructura eléctrica, dentro de una zona poblada o sobre un terreno que no admite el desarrollo. La bibliografía de selección de sitios solares combina información geográfica con criterios técnicos, ambientales y económicos, y advierte que los resultados dependen tanto de la calidad de las capas como de los pesos asignados [4]-[6].")
    add_p(doc, "En Santa Fe existen fuentes públicas para límites administrativos, asentamientos, redes eléctricas y estaciones transformadoras, pero presentan escalas, geometrías y coberturas diferentes [9]-[12]. BAHRA representa los asentamientos como puntos; por lo tanto, no entrega polígonos urbanos listos para excluir. La red de distribución y las estaciones transformadoras provienen de conjuntos distintos. La radiación solar se obtiene de un reanálisis climático y no de mediciones de campo. Integrar esas fuentes exige transformar sistemas de coordenadas, validar capas, normalizar métricas y conservar la trazabilidad de cada decisión.")
    add_p(doc, "El proyecto aborda ese problema con un producto mínimo viable reproducible. La provincia se divide en celdas; cada celda recibe métricas de radiación y proximidad a infraestructura; las celdas que intersectan la aproximación de área urbana quedan excluidas. El algoritmo genético trabaja sobre candidatos válidos y devuelve un ranking. La salida no pretende establecer la ubicación económicamente óptima, sino reducir el espacio de búsqueda y producir evidencia verificable para estudios posteriores.")

    add_heading(doc, "3 Problema de investigación", 1)
    add_p(doc, "¿Cómo integrar datos climáticos y geoespaciales heterogéneos en un algoritmo genético reproducible que permita identificar y ordenar, dentro de Santa Fe, celdas potencialmente favorables para un parque solar fotovoltaico de dos hectáreas, respetando la exclusión urbana y explicitando las limitaciones que impiden interpretar el resultado como una decisión definitiva de localización?")

    add_heading(doc, "4 Objetivos de la investigación", 1)
    add_heading(doc, "4 1 Objetivo general", 2)
    add_p(doc, "Desarrollar y verificar un prototipo de optimización geoespacial basado en algoritmos genéticos que produzca un ranking trazable de ubicaciones potencialmente favorables para parques solares fotovoltaicos en Santa Fe.")
    add_heading(doc, "4 2 Objetivos específicos", 2)
    add_numbered(doc, [
        "Integrar fuentes oficiales de límites, asentamientos e infraestructura eléctrica con radiación ERA5-Land ARCO.",
        "Construir una grilla configurable y calcular, para cada celda, radiación estimada y distancias a líneas eléctricas y estaciones transformadoras.",
        "Excluir del conjunto de búsqueda las celdas que intersecten áreas urbanas aproximadas mediante un buffer configurable.",
        "Definir una función de aptitud normalizada y configurable cuyos pesos sumen uno.",
        "Implementar selección por torneo, elitismo, cruce, mutación y un registro de mejores soluciones.",
        "Comparar el TOP 10 del algoritmo genético con el ranking exhaustivo de las celdas válidas.",
        "Documentar la arquitectura, los parámetros, las fuentes, las salidas y las limitaciones del prototipo.",
    ])

    add_heading(doc, "5 Hipótesis y alcance", 1)
    add_p(doc, "Hipótesis de trabajo. Si las métricas de radiación y proximidad a infraestructura se normalizan de manera consistente, las restricciones urbanas se aplican antes de la búsqueda y la ejecución utiliza una semilla fija, entonces un algoritmo genético puede recuperar de forma reproducible las celdas con mayor aptitud del conjunto analizado.")
    add_p(doc, "La hipótesis se evalúa en una corrida sobre Santa Fe con celdas de 9 km y pesos iniciales de 0,50 para radiación, 0,20 para proximidad a líneas y 0,30 para proximidad a transformadores. La superficie nominal de dos hectáreas se conserva como requisito de trazabilidad; el MVP no verifica continuidad ni disponibilidad física del terreno. El alcance es de preselección regional.")

    add_heading(doc, "6 Marco teórico", 1)
    add_heading(doc, "6 1 Selección espacial multicriterio", 2)
    add_p(doc, "Un problema de localización convierte alternativas geográficas en unidades comparables mediante criterios y restricciones. Los sistemas de información geográfica permiten superponer capas, medir distancias y excluir zonas; el análisis multicriterio combina esas evidencias en una regla de decisión [4]. En plantas solares, la irradiación, la proximidad a la red, las restricciones de uso del suelo y la accesibilidad aparecen de forma recurrente [5], [6]. El presente modelo adopta una suma ponderada de tres criterios activos y separa la exclusión urbana como una restricción dura.")
    add_heading(doc, "6 2 Algoritmos genéticos", 2)
    add_p(doc, "Los algoritmos genéticos son métodos de búsqueda poblacional inspirados en selección, recombinación y mutación [1]-[3]. Cada individuo codifica una solución; la función de aptitud cuantifica su calidad; la selección aumenta la presencia de individuos aptos; el cruce produce descendencia y la mutación mantiene diversidad. El elitismo conserva las mejores soluciones. Al usar una semilla pseudoaleatoria fija se obtiene repetibilidad, aunque no desaparece la naturaleza heurística del método.")
    add_p(doc, "En este prototipo, un individuo es un índice entero que referencia una fila del conjunto fijo de celdas válidas. La selección se realiza por torneos de tres individuos. El cruce mezcla dos índices y redondea el resultado; la mutación reemplaza el índice por otro candidato uniforme. Un registro de mejores soluciones conserva los individuos observados durante toda la corrida. Esta representación es compacta, pero el orden de las filas no es una variable geográfica continua; por eso el cruce aritmético no garantiza proximidad espacial ni semántica. Esta limitación se considera al interpretar los resultados.")
    add_heading(doc, "6 3 Función de aptitud", 2)
    add_p(doc, "Las métricas se normalizan al intervalo [0, 1]. Para radiación, el valor más alto recibe uno. Para distancias, la normalización se invierte para que la menor distancia reciba uno. Con pesos ws, wl y wt, la aptitud se calcula como:")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_repeat_font(p.add_run("fitness = ws · radiación + wl · proximidad a líneas + wt · proximidad a transformadores"), size=11.5, italic=True)
    add_p(doc, "Los pesos deben sumar uno. Un peso nulo desactiva el criterio y evita requerir su fuente. Esta decisión permite ensayar configuraciones, pero los valores empleados en la corrida son iniciales y no han sido calibrados mediante consulta a expertos, análisis de sensibilidad ni evidencia económica.")
    add_heading(doc, "6 4 Radiación solar y ERA5 Land", 2)
    add_p(doc, "ERA5-Land es un reanálisis terrestre global con resolución espacial aproximada de 9 km y muestreo horario [7]. El producto ARCO permite acceder a subconjuntos de variables sin descargar archivos completos y está organizado para consultas espaciales o temporales eficientes [8]. El prototipo usa la radiación solar descendente de superficie, convierte J/m² a kWh/m² y resume enero, abril, julio y octubre desde 1970 hasta el último mes representativo disponible. Cada media mensual histórica se lleva a una media diaria, se proyecta a los días de su estación y luego se suman las cuatro estaciones. El resultado es una estimación anual basada en meses representativos, no una suma observada de los doce meses.")
    add_heading(doc, "6 5 Calidad de datos y trazabilidad", 2)
    add_p(doc, "La validez del ranking depende de la correspondencia entre fuente, geometría, fecha y variable. El límite provincial proviene del IGN [9]; los asentamientos, de BAHRA [10]; las líneas y estaciones transformadoras, de la Secretaría de Energía [11], [12]. El sistema registra metadatos, usa caché, detiene el procesamiento cuando falta una capa activa y evita llamadas de red dentro del algoritmo genético. Las salidas incluyen la configuración, el período climático, los candidatos, el ranking y una advertencia de alcance.")

    add_heading(doc, "7 Concreción del modelo", 1)
    add_heading(doc, "7 1 Arquitectura", 2)
    add_picture(doc, "figura_arquitectura.png", 15.5, "Figura 1 Arquitectura funcional del prototipo")
    add_p(doc, "La separación entre adquisición, procesamiento y optimización evita que el algoritmo dependa de servicios externos durante la búsqueda. Los módulos de ingesta descargan y validan las capas; el preprocesamiento construye la grilla, transforma coordenadas, calcula distancias y genera métricas; el repositorio persiste los datos en SQLite; el módulo genético lee únicamente candidatos válidos; el reporte produce CSV, JSON y un mapa HTML.")
    add_heading(doc, "7 2 Fuentes y transformaciones", 2)
    rows = [
        ["Límite provincial", "IGN Unidades Territoriales", "Polígono", "Recorte de Santa Fe"],
        ["Áreas urbanas", "BAHRA", "Puntos", "Buffer de 3 km para LOCALIDAD"],
        ["Líneas eléctricas", "Secretaría de Energía", "Líneas", "Distancia al elemento más cercano"],
        ["Transformadores", "Secretaría de Energía", "Puntos", "Distancia al elemento más cercano"],
        ["Radiación", "Copernicus ERA5-Land ARCO", "Grilla climática", "Vecino más cercano y climatología"],
    ]
    add_table(doc, ["Capa", "Fuente", "Geometría", "Uso"], rows, widths=[3.2, 4.4, 2.6, 5.6], font_size=9)
    add_caption(doc, "Tabla 1 Fuentes y función dentro del modelo")
    add_heading(doc, "7 3 Flujo de procesamiento", 2)
    add_numbered(doc, [
        "Validar la configuración, la región Santa Fe y la coherencia de los pesos.",
        "Descargar o recuperar de caché el límite, BAHRA, líneas y transformadores.",
        "Construir una grilla de 9 km en un sistema proyectado métrico.",
        "Bufferizar 3 km los puntos BAHRA de tipo LOCALIDAD y excluir celdas que intersectan esos buffers.",
        "Calcular la distancia mínima desde cada celda a líneas y transformadores.",
        "Consultar ERA5-Land ARCO por bloques espaciales y resumir 227 meses representativos.",
        "Normalizar las métricas, persistir candidatos y ejecutar el algoritmo genético.",
        "Generar ranking, metadatos de corrida, candidatos y mapa interactivo.",
    ])
    add_heading(doc, "7 4 Especificación técnica", 2)
    add_table(doc, ["Componente", "Especificación"], [
        ["Lenguaje", "Python 3.11 o superior; corrida verificada con Python 3.14 en Windows"],
        ["Datos y cálculo", "pandas, NumPy, xarray y Zarr"],
        ["Geoespacial", "GeoPandas, Shapely, pyproj y pyogrio"],
        ["Persistencia", "SQLite mediante SQLAlchemy"],
        ["Visualización", "Folium y mosaicos OpenStreetMap con atribución"],
        ["Configuración", "YAML validado con Pydantic y credencial ARCO en archivo .env"],
        ["Infraestructura", "Equipo Windows con conexión para la primera ingesta y almacenamiento local para caché y base SQLite"],
    ], widths=[4.2, 11.6], font_size=9.5)
    add_caption(doc, "Tabla 2 Software e infraestructura del prototipo")

    add_heading(doc, "8 Diseño experimental", 1)
    add_p(doc, f"La corrida registrada como {RUN['metadata']['run_id']} utilizó celdas de 9 km, un parque nominal de 2 ha, período climático 1970-01 a 2026-07 y los meses 1, 4, 7 y 10. La estimación se construyó con 227 meses. El algoritmo usó 50 individuos, 100 generaciones, probabilidad de cruce 0,75, probabilidad de mutación 0,05, elitismo de dos, torneos de tres y semilla 42.")
    add_table(doc, ["Parámetro", "Valor"], [
        ["Pesos del fitness", "Radiación 0,50; líneas 0,20; transformadores 0,30"],
        ["Población", "50"], ["Generaciones", "100"], ["Cruce", "0,75"],
        ["Mutación", "0,05"], ["Elitismo", "2"], ["Torneo", "3"], ["Semilla", "42"],
    ], widths=[6.0, 9.8], font_size=10)
    add_caption(doc, "Tabla 3 Configuración de la corrida analizada")
    add_p(doc, f"La verificación tuvo tres niveles. Primero, se revisaron los archivos de salida y el log de ejecución. Segundo, se ejecutaron 11 pruebas automatizadas sobre configuración flexible, desactivación de criterios, caché, persistencia y generación del mapa; todas finalizaron correctamente. Tercero, se calculó el ranking exhaustivo de las {RUN['valid']} celdas válidas con la misma función de aptitud y se comparó con el TOP 10 genético.")

    add_heading(doc, "9 Resultados", 1)
    add_p(doc, f"La grilla produjo {RUN['total']} celdas. El criterio urbano excluyó {RUN['excluded']}, equivalentes al {RUN['excluded_pct']:.1f} % del total; no se registraron exclusiones por falta de datos climáticos. El algoritmo consideró {RUN['valid']} candidatos. La Figura 2 muestra la distribución espacial de las celdas y el TOP 10.")
    add_picture(doc, "figura_mapa_candidatos.png", 11.2, "Figura 2 Celdas evaluadas y posiciones del ranking TOP 10")
    r = RUN["ranking"]
    table_rows = []
    for x in r.itertuples():
        table_rows.append([
            str(int(x.rank)), str(int(x.grid_cell_id)), f"{x.latitude:.4f}", f"{x.longitude:.4f}",
            f"{x.fitness:.4f}", f"{x.solar_annual_kwh_m2:.1f}", f"{x.distance_to_power_line_km:.1f}", f"{x.distance_to_transformer_km:.1f}",
        ])
    add_table(doc, ["Rango", "Celda", "Lat", "Lon", "Fitness", "kWh/m²/año", "Línea km", "Trafo km"], table_rows, widths=[1.2, 1.4, 2.0, 2.0, 1.7, 2.4, 2.0, 2.0], font_size=8.2)
    add_caption(doc, "Tabla 4 Ranking de ubicaciones potencialmente favorables")
    add_p(doc, f"La celda {int(RUN['first'].grid_cell_id)} obtuvo fitness {RUN['first'].fitness:.4f}, radiación anual estimada de {RUN['first'].solar_annual_kwh_m2:.1f} kWh/m²/año y distancia nula a la geometría de una línea eléctrica. Su distancia a la estación transformadora más cercana fue {RUN['first'].distance_to_transformer_km:.1f} km. Las diez soluciones genéticas presentaron una radiación anual estimada entre {RUN['solar_min']:.1f} y {RUN['solar_max']:.1f} kWh/m²/año; las {RUN['zero_lines']} registraron 0 km a una geometría de línea.")
    add_picture(doc, "figura_componentes_fitness.png", 15.5, "Figura 3 Contribución ponderada de cada criterio al fitness")
    add_p(doc, f"La ejecución genética no recuperó el TOP 10 exhaustivo de los {RUN['valid']} candidatos. Compartió {RUN['overlap']} celdas con ese conjunto y omitió la mejor alternativa de referencia: la celda {int(RUN['exhaustive_best'].grid_cell_id)}, cuyo fitness exhaustivo fue {RUN['exhaustive_best'].fitness:.4f}. La diferencia indica que los parámetros y la representación actuales no exploran de manera suficiente el espacio ampliado por la grilla de 9 km.")
    add_picture(doc, "figura_convergencia.png", 15.5, "Figura 4 Evolución del mejor fitness y del fitness medio")

    add_heading(doc, "10 Discusión", 1)
    add_p(doc, f"La corrida es reproducible con la semilla y los parámetros fijados, pero no verifica la hipótesis de recuperación del mejor ranking: el TOP 10 genético solo comparte {RUN['overlap']} celdas con el exhaustivo. La grilla de 9 km amplió el espacio de búsqueda a {RUN['valid']} alternativas y expuso una limitación que no aparecía en la corrida de 30 km. La búsqueda completa sigue siendo viable para este tamaño y funciona como referencia necesaria.")
    add_p(doc, "La función de aptitud deja visible el compromiso entre recurso solar e infraestructura. La primera posición no es la más cercana a transformadores; gana por radiación y proximidad a líneas. Esto confirma que un único valor de fitness puede ocultar compensaciones relevantes. Antes de usar el ranking en una decisión real deben ejecutarse análisis de sensibilidad de pesos y, de ser posible, formular el problema como optimización multiobjetivo.")
    add_p(doc, "La representación genética requiere una revisión. El índice de fila es una categoría, aunque el cruce lo trata como una magnitud continua. Dos índices próximos pueden corresponder a celdas lejanas y viceversa. Una versión futura debería usar coordenadas, identificadores organizados espacialmente, operadores categóricos o una formulación de conjunto cuando se busquen varias instalaciones. También conviene comparar múltiples semillas y medir diversidad, convergencia y estabilidad del ranking.")
    add_p(doc, "La exclusión urbana es una aproximación. BAHRA aporta puntos y el buffer uniforme de 3 km no reproduce la extensión real de cada localidad. La grilla de 9 km sigue siendo mucho mayor que un parque de 2 ha y una celda válida puede contener restricciones no modeladas. Además, la radiación estimada con cuatro meses representativos no equivale a una simulación fotovoltaica completa. Estas limitaciones impiden convertir el ranking en una recomendación de obra.")

    add_heading(doc, "11 Conclusiones", 1)
    add_p(doc, f"El proyecto concretó un prototipo funcional que integra datos oficiales, procesamiento geoespacial, persistencia local, optimización genética y productos de salida trazables. La corrida analizada identificó {RUN['valid']} celdas válidas entre {RUN['total']} y situó a la celda {int(RUN['first'].grid_cell_id)} en primer lugar del ranking genético con fitness {RUN['first'].fitness:.4f}. Las 11 pruebas automatizadas finalizaron sin errores, pero el TOP 10 no coincidió con el ranking exhaustivo.")
    add_p(doc, "El resultado cumple el objetivo de preselección y deja una base de software para continuar la investigación. No prueba que las ubicaciones sean técnica, ambiental, legal o económicamente viables. La conclusión defendible es más acotada: el sistema ordena de forma reproducible celdas potencialmente favorables según las variables y pesos incorporados.")

    add_heading(doc, "12 Líneas de continuidad", 1)
    add_bullets(doc, [
        "Reemplazar los buffers de puntos urbanos por polígonos actualizados de ocupación del suelo.",
        "Mantener fuera del modelo las variables orográficas y de nubosidad en esta etapa y documentar explícitamente ese alcance.",
        "Agregar capacidad de líneas y subestaciones, costos de conexión, costo del terreno y restricciones regulatorias.",
        "Ejecutar sensibilidad de pesos y una formulación multiobjetivo que exponga soluciones de compromiso.",
        "Rediseñar la codificación genética y comparar el algoritmo con búsqueda exhaustiva, recocido simulado y métodos voraces.",
        "Validar las ubicaciones priorizadas con cartografía de mayor resolución y relevamientos de campo.",
    ])

    add_heading(doc, "Referencias", 1)
    for ref in REFERENCES:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(.6)
        p.paragraph_format.first_line_indent = Cm(-.6)
        p.paragraph_format.space_after = Pt(4)
        set_repeat_font(p.add_run(ref), size=9.5)

    add_heading(doc, "Anexo técnico", 1)
    add_heading(doc, "A Configuración reproducible", 2)
    add_p(doc, "Comando principal en Windows: .\\.venv\\Scripts\\python.exe -m src.main --run-all")
    add_p(doc, "Comando de pruebas: .\\.venv\\Scripts\\python.exe -m unittest discover -s tests -v")
    add_heading(doc, "B Productos de la corrida", 2)
    add_bullets(doc, [
        "results/ranking.csv con el TOP 10 y el desglose climático.",
        f"results/candidate_locations.csv con las {RUN['total']} celdas y su estado.",
        "results/optimization_run.json con parámetros, período y advertencia de alcance.",
        "results/map.html con capas y ubicaciones priorizadas.",
        "data/processed/parque_solar.sqlite con datos y resultados persistidos.",
    ])

    doc.core_properties.title = "Optimización de ubicaciones de parques solares mediante algoritmos genéticos"
    doc.core_properties.subject = "Documento guía de investigación y concreción del modelo"
    doc.core_properties.author = ""
    doc.save(PROFILE)


def set_columns(section, num=2, space_twips=567):
    sect_pr = section._sectPr
    cols = sect_pr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sect_pr.append(cols)
    cols.set(qn("w:num"), str(num))
    cols.set(qn("w:space"), str(space_twips))


def article_document() -> None:
    doc = Document()
    configure_base(doc, body_size=12, line=1.0, after=3)
    sec = doc.sections[0]
    sec.top_margin = Cm(2.5)
    sec.bottom_margin = Cm(2.5)
    sec.left_margin = Cm(2.5)
    sec.right_margin = Cm(2.5)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    set_repeat_font(p.add_run("Optimización de ubicaciones de parques solares mediante algoritmos genéticos"), size=16, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    set_repeat_font(p.add_run(" "), size=14, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    set_repeat_font(p.add_run("Universidad Tecnológica Nacional Facultad Regional Buenos Aires"), size=12, bold=True, italic=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    set_repeat_font(p.add_run("Versión para evaluación ciega"), size=10, italic=True)

    sec2 = doc.add_section(WD_SECTION.CONTINUOUS)
    sec2.top_margin = Cm(2.5)
    sec2.bottom_margin = Cm(2.5)
    sec2.left_margin = Cm(2.5)
    sec2.right_margin = Cm(2.5)
    set_columns(sec2, 2, 567)

    def h(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(7)
        p.paragraph_format.space_after = Pt(3)
        set_repeat_font(p.add_run(text), size=12, bold=True)
        set_keep(p, keep_next=True)
        return p

    def body(text, size=12, italic=False):
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(3)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        set_repeat_font(p.add_run(text), size=size, italic=italic)
        return p

    h("Abstract")
    body(f"Este trabajo presenta un prototipo reproducible para identificar celdas potencialmente favorables para un parque solar fotovoltaico en Santa Fe. El sistema integra radiación ERA5-Land ARCO, distancias a líneas eléctricas y estaciones transformadoras, y una exclusión urbana construida a partir de puntos BAHRA. La corrida analizada generó {RUN['total']} celdas de 9 km; {RUN['excluded']} fueron excluidas y {RUN['valid']} ingresaron a la optimización. La celda {int(RUN['first'].grid_cell_id)} encabezó el ranking genético con fitness {RUN['first'].fitness:.4f} y radiación anual estimada de {RUN['first'].solar_annual_kwh_m2:.1f} kWh/m². El TOP 10 genético compartió {RUN['overlap']} celdas con el exhaustivo y omitió su mejor alternativa. Las 11 pruebas automatizadas finalizaron correctamente. El resultado confirma la ejecución reproducible, pero muestra que la configuración genética actual no garantiza recuperar las mejores alternativas al aumentar la resolución.", size=10, italic=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    set_repeat_font(p.add_run("Palabras Clave "), size=10, bold=True)
    set_repeat_font(p.add_run("algoritmos genéticos, energía solar, análisis geoespacial, selección de sitios, ERA5-Land"), size=10)

    h("Introducción")
    body("La localización preliminar de una planta fotovoltaica exige comparar alternativas espaciales según recurso solar, infraestructura, restricciones territoriales y costos. Los enfoques GIS-MCDA organizan esas evidencias, pero el resultado depende de las capas, las reglas de exclusión y los pesos [4]-[6]. Este proyecto estudia si un algoritmo genético puede recuperar de manera reproducible las mejores celdas de una grilla provincial a partir de tres criterios cuantitativos y una restricción urbana.")
    body("La contribución es un flujo completo y trazable. La ingesta usa fuentes oficiales argentinas para límite, asentamientos y red eléctrica [9]-[12], y ERA5-Land ARCO para radiación [7], [8]. El algoritmo no consulta servicios externos: opera sobre métricas persistidas en una base local. La salida es un ranking de preselección, no una decisión de inversión.")

    h("Elementos del trabajo y metodología")
    body(f"La provincia de Santa Fe se discretizó mediante {RUN['total']} celdas de 9 km en un sistema de coordenadas proyectado. BAHRA describe asentamientos como puntos; se seleccionaron registros de tipo LOCALIDAD y se aplicó un buffer de 3 km. Toda celda que intersectó un buffer se marcó como inválida. Para las restantes se calculó la distancia mínima a líneas eléctricas y estaciones transformadoras.")
    body("ERA5-Land ofrece reanálisis horario terrestre con resolución aproximada de 9 km [7]. El acceso ARCO permite leer subconjuntos sin descargar archivos completos [8]. Para cada centroide se eligió el punto climático más cercano y se resumieron enero, abril, julio y octubre desde 1970 hasta 2026. Las medias mensuales se convirtieron a irradiación diaria, se proyectaron a los días de cada estación y se sumaron como estimación anual.")
    body("Las tres métricas se normalizaron entre cero y uno. La radiación conserva el sentido creciente; las distancias se invierten para representar proximidad. La función utilizada fue fitness = 0,50·radiación + 0,20·proximidad a líneas + 0,30·proximidad a transformadores. Los pesos son iniciales y no están científicamente calibrados.")
    body("Cada individuo del algoritmo representa el índice de una celda válida. La población fue de 50 individuos durante 100 generaciones, con selección por torneo de tamaño tres, elitismo de dos, cruce con probabilidad 0,75, mutación con probabilidad 0,05 y semilla 42. Un registro histórico conservó las mejores soluciones vistas. La implementación sigue los componentes habituales de los algoritmos genéticos [1]-[3].")
    body(f"La verificación incluyó revisión de metadatos, 11 pruebas automatizadas y una comparación contra el ordenamiento exhaustivo de las {RUN['valid']} celdas válidas.")

    h("Resultados")
    body(f"El procesamiento excluyó {RUN['excluded']} celdas por intersección urbana, {RUN['excluded_pct']:.1f} % del total, y no registró faltantes climáticos. La Figura 1 ubica las diez soluciones genéticas mejor clasificadas. La primera fue la celda {int(RUN['first'].grid_cell_id)}, en latitud {RUN['first'].latitude:.4f} y longitud {RUN['first'].longitude:.4f}, con fitness {RUN['first'].fitness:.4f}. Su radiación anual estimada fue {RUN['first'].solar_annual_kwh_m2:.1f} kWh/m²; la distancia a una línea fue {RUN['first'].distance_to_power_line_km:.1f} km y la distancia al transformador más cercano, {RUN['first'].distance_to_transformer_km:.1f} km.")
    add_picture(doc, "figura_mapa_candidatos.png", 7.1, "Figura 1 Celdas y ubicaciones TOP 10")
    rows = []
    r = RUN["ranking"]
    for x in r.itertuples():
        rows.append([str(int(x.rank)), str(int(x.grid_cell_id)), f"{x.fitness:.3f}", f"{x.solar_annual_kwh_m2:.0f}", f"{x.distance_to_transformer_km:.1f}"])
    add_table(doc, ["R", "Celda", "Fit", "Solar", "Trafo"], rows, widths=[.7, 1.1, 1.1, 1.4, 1.4], font_size=8)
    add_caption(doc, "Tabla 1 Ranking y métricas principales")
    body(f"Las diez celdas presentaron entre {RUN['solar_min']:.1f} y {RUN['solar_max']:.1f} kWh/m²/año. Las {RUN['zero_lines']} registraron distancia nula a una geometría de línea. El TOP 10 genético compartió {RUN['overlap']} celdas con el exhaustivo y no incluyó la mejor alternativa de referencia, la celda {int(RUN['exhaustive_best'].grid_cell_id)}. El conjunto de pruebas finalizó sin errores.")
    add_picture(doc, "figura_componentes_fitness.png", 7.1, "Figura 2 Aportes ponderados al fitness")

    h("Discusión")
    body(f"La comparación exhaustiva muestra que la configuración genética no alcanzó el ranking de referencia: solo {RUN['overlap']} de sus diez soluciones pertenecen al TOP 10 exhaustivo. La mayor resolución hizo visible una limitación de exploración que debe corregirse antes de interpretar el ranking genético como selección final.")
    body(f"El primer puesto genético, la celda {int(RUN['first'].grid_cell_id)}, combina el máximo puntaje solar con distancia nula a líneas, aunque permanece lejos de una estación transformadora. La mejor alternativa exhaustiva fue la celda {int(RUN['exhaustive_best'].grid_cell_id)}, con fitness {RUN['exhaustive_best'].fitness:.4f}.")
    body("La codificación también limita la interpretación. El índice de fila es categórico, pero el cruce aritmético supone continuidad. Índices cercanos no necesariamente representan sitios cercanos. Una extensión debería emplear una codificación espacial o un operador categórico, ejecutar múltiples semillas y comparar estabilidad y diversidad.")
    body("La principal amenaza externa sigue siendo la resolución y representación de los datos. Un buffer uniforme alrededor de puntos BAHRA no representa límites urbanos reales. Una celda de 9 km tampoco demuestra que haya 20.000 m² contiguos utilizables. La estimación climática usa cuatro meses representativos y no modela el rendimiento eléctrico de un sistema fotovoltaico. Por estas razones, el ranking solo orienta estudios posteriores.")

    h("Conclusión")
    body(f"El prototipo integró fuentes oficiales, métricas geoespaciales, persistencia y optimización genética en un proceso reproducible. Identificó {RUN['valid']} candidatos válidos entre {RUN['total']} celdas. La corrida no permite aceptar la hipótesis de recuperación de las mejores alternativas, porque el TOP 10 genético no coincidió con el exhaustivo.")
    body("El resultado no establece viabilidad técnica ni económica. El siguiente avance metodológico debe concentrarse en revisar la representación genética, la exploración, el tamaño poblacional y la estabilidad entre semillas, manteniendo fuera del alcance actual las variables orográficas y de nubosidad.")

    h("Referencias")
    for ref in REFERENCES:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(.35)
        p.paragraph_format.first_line_indent = Cm(-.35)
        p.paragraph_format.space_after = Pt(2)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_repeat_font(p.add_run(ref), size=8.5)

    h("Datos de contacto")
    body("Universidad Tecnológica Nacional Facultad Regional Buenos Aires", size=10, italic=True)
    doc.core_properties.title = "Optimización de ubicaciones de parques solares mediante algoritmos genéticos"
    doc.core_properties.subject = "Artículo de cátedra"
    doc.core_properties.author = ""
    doc.save(ARTICLE)


if __name__ == "__main__":
    profile_document()
    article_document()
    print(PROFILE)
    print(ARTICLE)
